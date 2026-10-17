from pathlib import Path
import hashlib
import re
import pymupdf as fitz
from pptx import Presentation


SUPPORTED_EXTENSIONS = {".pdf", ".pptx"}
SECTION_HEADING = re.compile(
    r"^\s*\d+\s*-\s*\d+(?:[.-]\d+)*(?:[.)])?\s+\S"
)
NUMBERED_STEP = re.compile(r"^\s*(?:\d+[.)]|[A-Z][.)])\s+\S")


def extract_text(file_path: Path) -> list[dict]:
    """Extract text while preserving page/slide references."""
    pages = []

    if file_path.suffix.lower() == ".pdf":
        with fitz.open(file_path) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                blocks = page.get_text("blocks", sort=True)
                text = "\n\n".join(
                    block[4].strip()
                    for block in blocks
                    if len(block) > 6
                    and block[6] == 0
                    and block[4].strip()
                )
                if text:
                    pages.append({
                        "location": f"Page {page_number}",
                        "text": text,
                    })

    elif file_path.suffix.lower() == ".pptx":
        presentation = Presentation(file_path)

        for slide_number, slide in enumerate(
            presentation.slides, start=1
        ):
            parts = []

            for shape in sorted(
                slide.shapes,
                key=lambda item: (item.top, item.left),
            ):
                if shape.has_text_frame and shape.text.strip():
                    parts.append(shape.text.strip())

                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [
                            cell.text.strip() for cell in row.cells
                        ]
                        parts.append(" | ".join(cells))

            text = "\n".join(parts).strip()
            if text:
                pages.append({
                    "location": f"Slide {slide_number}",
                    "text": text,
                })

    return pages


def chunk_text(text: str, chunk_size: int = 1000,
               overlap: int = 150) -> list[str]:
    """Split text into overlapping character chunks."""
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Invalid chunk_size or overlap")

    words = text.split()
    chunks = []
    current_words = []
    current_length = 0

    for word in words:
        added_length = len(word) + (1 if current_words else 0)

        if current_words and current_length + added_length > chunk_size:
            chunks.append(" ".join(current_words))

            # Retain trailing words for overlap.
            tail = []
            tail_length = 0

            for previous_word in reversed(current_words):
                length = len(previous_word) + (
                    1 if tail else 0
                )
                if tail and tail_length + length > overlap:
                    break
                tail.insert(0, previous_word)
                tail_length += length

            current_words = tail
            current_length = len(" ".join(current_words))

        current_words.append(word)
        current_length += len(word) + (
            1 if len(current_words) > 1 else 0
        )

    if current_words:
        chunks.append(" ".join(current_words))

    return chunks


def chunk_page(
    text: str,
    chunk_size: int = 1000,
    overlap: int = 150,
    initial_section: str = "",
) -> list[dict[str, str]]:
    """Chunk a page while retaining section headings and step boundaries."""
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Invalid chunk_size or overlap")

    paragraphs = []
    current_lines = []

    for line in text.splitlines():
        line = line.strip()
        if not line:
            if current_lines:
                paragraphs.append(" ".join(current_lines))
                current_lines = []
            continue

        if (
            current_lines
            and (
                NUMBERED_STEP.match(line)
                or SECTION_HEADING.match(line)
            )
        ):
            paragraphs.append(" ".join(current_lines))
            current_lines = []
        current_lines.append(line)

    if current_lines:
        paragraphs.append(" ".join(current_lines))

    chunks = []
    section = initial_section
    current_parts = []
    current_length = 0

    def flush():
        if current_parts:
            chunks.append({
                "content": "\n".join(current_parts),
                "section": section,
            })

    for paragraph in paragraphs:
        if SECTION_HEADING.match(paragraph):
            flush()
            current_parts = [paragraph]
            current_length = len(paragraph)
            section = paragraph
            continue

        parts = [paragraph]
        if len(paragraph) > chunk_size:
            parts = chunk_text(paragraph, chunk_size, overlap)

        for part in parts:
            if not current_parts and section:
                current_parts = [section]
                current_length = len(section)

            added_length = len(part) + (1 if current_parts else 0)
            if current_parts and (
                current_length + added_length > chunk_size
            ) and not (section and current_parts == [section]):
                flush()
                current_parts = [section] if section else []
                current_length = len(section)

            if current_parts and current_parts[-1] == section:
                current_parts.append(part)
                current_length += len(part) + 1
            else:
                current_parts.append(part)
                current_length += len(part) + (
                    1 if len(current_parts) > 1 else 0
                )

    flush()
    return chunks


def stable_document_id(
    file_path: Path,
    base_dir: Path | None = None,
) -> str:
    """Create a stable identifier for a file."""
    resolved_path = file_path.resolve()
    if base_dir is not None:
        try:
            identity = resolved_path.relative_to(base_dir.resolve()).as_posix()
        except ValueError:
            identity = resolved_path.as_posix()
    else:
        identity = resolved_path.as_posix()
    return hashlib.sha256(identity.encode("utf-8")).hexdigest()